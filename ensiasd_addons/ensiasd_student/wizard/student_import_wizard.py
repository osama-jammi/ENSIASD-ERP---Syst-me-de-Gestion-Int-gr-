# -*- coding: utf-8 -*-
import base64
import io
import csv
import logging
from odoo import models, fields, api
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)


class StudentImportWizard(models.TransientModel):
    """
    Wizard d'import en masse des étudiants depuis un fichier CSV ou Excel.

    Flux :
      1. L'utilisateur choisit la filière, le niveau, l'année et uploade le fichier.
      2. Il clique sur « Aperçu » pour valider les données sans les enregistrer.
      3. Il clique sur « Importer » pour créer les fiches étudiants.

    Colonnes attendues dans le fichier (ordre configurable) :
      nom_complet | cne | cin | email | telephone | mobile | date_naissance |
      lieu_naissance | sexe | nationalite | adresse | ville |
      groupe | contact_urgence_nom | contact_urgence_tel

    Seules les colonnes nom_complet et cne sont obligatoires.
    """
    _name = 'ensiasd.student.import.wizard'
    _description = "Assistant d'import en masse des étudiants"

    # ------------------------------------------------------------------
    # Paramètres globaux de l'import
    # ------------------------------------------------------------------

    filiere_id = fields.Many2one(
        'ensiasd.filiere',
        string='Filière',
        required=True,
        help="Filière dans laquelle inscrire tous les étudiants du fichier"
    )

    niveau = fields.Selection([
        ('1', '1ère année'),
        ('2', '2ème année'),
        ('3', '3ème année'),
    ], string='Niveau', required=True, default='1')

    annee_id = fields.Many2one(
        'ensiasd.annee',
        string='Année académique',
        required=True,
        default=lambda self: self._default_annee()
    )

    groupe_id = fields.Many2one(
        'ensiasd.groupe',
        string='Groupe par défaut',
        help="Groupe affecté à tous les étudiants (peut être surchargé colonne par colonne)"
    )

    state_after_import = fields.Selection([
        ('draft', 'Brouillon'),
        ('inscrit', 'Inscrit'),
    ], string='État après import', default='inscrit',
        help="État des fiches créées. 'Inscrit' déclenche l'inscription automatique aux modules."
    )

    auto_inscribe_modules = fields.Boolean(
        string='Inscription automatique aux modules',
        default=True,
        help="Si coché, inscrit automatiquement chaque étudiant aux modules obligatoires de la filière."
    )

    # ------------------------------------------------------------------
    # Fichier
    # ------------------------------------------------------------------

    file = fields.Binary(string='Fichier CSV / Excel', required=True, attachment=False)
    filename = fields.Char(string='Nom du fichier')

    file_type = fields.Selection([
        ('csv', 'CSV (.csv)'),
        ('xlsx', 'Excel (.xlsx)'),
    ], string='Format', default='csv', required=True,
        compute='_compute_file_type', store=True, readonly=False
    )

    delimiter = fields.Selection([
        (';', 'Point-virgule  ;'),
        (',', 'Virgule  ,'),
        ('\t', 'Tabulation'),
    ], string='Délimiteur CSV', default=';')

    encoding = fields.Selection([
        ('utf-8-sig', 'UTF-8 (avec BOM)'),
        ('utf-8', 'UTF-8'),
        ('latin-1', 'Latin-1 / ISO-8859-1'),
        ('windows-1252', 'Windows-1252'),
    ], string='Encodage CSV', default='utf-8-sig')

    skip_header = fields.Boolean(
        string='Ignorer la 1ère ligne (en-tête)',
        default=True
    )

    # ------------------------------------------------------------------
    # Mapping des colonnes (numéro 1-based, 0 = colonne absente)
    # ------------------------------------------------------------------

    col_nom = fields.Integer('Colonne Nom complet', default=1)
    col_cne = fields.Integer('Colonne CNE', default=2)
    col_cin = fields.Integer('Colonne CIN', default=3)
    col_email = fields.Integer('Colonne Email', default=4)
    col_telephone = fields.Integer('Colonne Téléphone', default=5)
    col_mobile = fields.Integer('Colonne Mobile', default=6)
    col_date_naissance = fields.Integer(
        'Colonne Date naissance',
        default=7,
        help="Format attendu : YYYY-MM-DD ou DD/MM/YYYY"
    )
    col_lieu_naissance = fields.Integer('Colonne Lieu naissance', default=8)
    col_sexe = fields.Integer(
        'Colonne Sexe', default=9,
        help="Valeurs acceptées : M/m/masculin/Masculin → male ; F/f/féminin/Féminin → female"
    )
    col_nationalite = fields.Integer('Colonne Nationalité', default=10)
    col_adresse = fields.Integer('Colonne Adresse', default=11)
    col_ville = fields.Integer('Colonne Ville', default=12)
    col_groupe = fields.Integer(
        'Colonne Groupe', default=0,
        help="Laisser à 0 pour utiliser le groupe par défaut ci-dessus"
    )
    col_contact_urgence_nom = fields.Integer('Colonne Contact urgence (nom)', default=0)
    col_contact_urgence_tel = fields.Integer('Colonne Contact urgence (tél)', default=0)

    # ------------------------------------------------------------------
    # Résultats
    # ------------------------------------------------------------------

    state = fields.Selection([
        ('draft', 'Configuration'),
        ('preview', 'Aperçu'),
        ('done', 'Terminé'),
    ], string='Étape', default='draft')

    preview_html = fields.Html(
        string='Aperçu des données',
        readonly=True
    )

    import_count = fields.Integer('Étudiants créés', readonly=True, default=0)
    update_count = fields.Integer('Étudiants mis à jour (CNE existant)', readonly=True, default=0)
    skip_count = fields.Integer('Lignes ignorées', readonly=True, default=0)
    error_count = fields.Integer('Erreurs', readonly=True, default=0)
    error_log = fields.Text('Détail des erreurs', readonly=True)
    result_summary = fields.Char('Résumé', readonly=True)

    # ------------------------------------------------------------------
    # Calculs automatiques
    # ------------------------------------------------------------------

    @api.model
    def _default_annee(self):
        config = self.env['ensiasd.config'].sudo().search([], limit=1)
        if config and config.annee_courante_id:
            return config.annee_courante_id.id
        annee = self.env['ensiasd.annee'].search([('state', '=', 'en_cours')], limit=1)
        return annee.id if annee else False

    @api.depends('filename')
    def _compute_file_type(self):
        for rec in self:
            if rec.filename:
                ext = rec.filename.rsplit('.', 1)[-1].lower()
                if ext in ('xlsx', 'xls'):
                    rec.file_type = 'xlsx'
                else:
                    rec.file_type = 'csv'
            else:
                rec.file_type = 'csv'

    # ------------------------------------------------------------------
    # Parsing
    # ------------------------------------------------------------------

    def _parse_file(self):
        """Décode et parse le fichier uploadé. Retourne list[list[str]]."""
        if not self.file:
            raise UserError("Veuillez sélectionner un fichier.")

        file_content = base64.b64decode(self.file)

        if self.file_type == 'csv':
            return self._parse_csv(file_content)
        else:
            return self._parse_xlsx(file_content)

    def _parse_csv(self, content_bytes):
        """Parse CSV et retourne list[list[str]]."""
        try:
            text = content_bytes.decode(self.encoding)
        except (UnicodeDecodeError, LookupError) as e:
            raise UserError(
                f"Impossible de lire le fichier avec l'encodage '{self.encoding}'.\n"
                f"Essayez un autre encodage. Détail : {e}"
            )

        reader = csv.reader(io.StringIO(text), delimiter=self.delimiter)
        rows = list(reader)

        if self.skip_header and rows:
            rows = rows[1:]

        return [row for row in rows if any(cell.strip() for cell in row)]

    def _parse_xlsx(self, content_bytes):
        """Parse XLSX et retourne list[list[str]]."""
        try:
            import openpyxl
        except ImportError:
            raise UserError(
                "Le module Python 'openpyxl' n'est pas installé.\n"
                "Lancez : pip install openpyxl  puis redémarrez Odoo."
            )

        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
        ws = wb.active

        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if self.skip_header and i == 1:
                continue
            row_str = [str(cell) if cell is not None else '' for cell in row]
            if any(cell.strip() for cell in row_str):
                rows.append(row_str)

        wb.close()
        return rows

    # ------------------------------------------------------------------
    # Helpers d'extraction
    # ------------------------------------------------------------------

    @staticmethod
    def _col(row, col_num):
        """Retourne la valeur de la colonne (1-based). '' si hors limites ou col=0."""
        if col_num <= 0 or col_num > len(row):
            return ''
        return str(row[col_num - 1]).strip()

    @staticmethod
    def _parse_date(val):
        """Convertit DD/MM/YYYY ou YYYY-MM-DD en date ou None."""
        val = val.strip()
        if not val or val in ('-', 'None', 'nan'):
            return None
        for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
            try:
                from datetime import datetime
                return datetime.strptime(val, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_sexe(val):
        """Normalise la valeur de sexe vers 'male'/'female'/None."""
        v = val.strip().lower()
        if v in ('m', 'h', 'male', 'masculin', 'homme'):
            return 'male'
        if v in ('f', 'female', 'féminin', 'feminin', 'femme'):
            return 'female'
        return None

    def _group_vals(self, groupe_name):
        """Valeurs de création d'un groupe : année (obligatoire) + filière/niveau si le modèle les porte."""
        Groupe = self.env['ensiasd.groupe']
        vals = {'name': groupe_name}
        if 'code' in Groupe._fields:
            # Le code est obligatoire sur ensiasd.groupe : on reprend le nom (ex. G1, G-TEST)
            vals['code'] = groupe_name.strip().upper()
        if 'annee_id' in Groupe._fields:
            vals['annee_id'] = self.annee_id.id
        if 'filiere_id' in Groupe._fields:
            vals['filiere_id'] = self.filiere_id.id
        if 'niveau' in Groupe._fields:
            try:
                allowed = dict(Groupe._fields['niveau']._description_selection(self.env))
            except Exception:
                allowed = {}
            if self.niveau in allowed:
                vals['niveau'] = self.niveau
        return vals

    def _resolve_groupe(self, groupe_name, dry_run=False):
        """Cherche un groupe par nom (dans l'année choisie), crée-le si absent (sauf en dry_run)."""
        if not groupe_name:
            return self.groupe_id.id if self.groupe_id else False
        Groupe = self.env['ensiasd.groupe']
        domain = [('name', '=ilike', groupe_name)]
        if 'annee_id' in Groupe._fields:
            domain.append(('annee_id', '=', self.annee_id.id))
        groupe = Groupe.search(domain, limit=1)
        if groupe:
            return groupe.id
        if dry_run:
            # Aperçu : ne rien écrire en base
            return False
        return Groupe.create(self._group_vals(groupe_name)).id

    def _build_student_vals(self, row, dry_run=False):
        """
        Construit le dictionnaire de valeurs pour create/write à partir d'une ligne.
        Retourne (vals_dict, erreurs_list).
        """
        errors = []
        col = self._col

        nom = col(row, self.col_nom)
        cne = col(row, self.col_cne)

        if not nom:
            errors.append("Nom complet manquant")
        if not cne:
            errors.append("CNE manquant")

        if errors:
            return None, errors

        vals = {
            'name': nom,
            'cne': cne,
            'filiere_id': self.filiere_id.id,
            'niveau': self.niveau,
            'annee_inscription': self.annee_id.id,
            'annee_courante_id': self.annee_id.id,
        }

        cin = col(row, self.col_cin)
        if cin:
            vals['cin'] = cin

        email = col(row, self.col_email)
        if email:
            vals['email'] = email

        phone = col(row, self.col_telephone)
        if phone:
            vals['phone'] = phone

        mobile = col(row, self.col_mobile)
        if mobile:
            vals['mobile'] = mobile

        date_naissance_str = col(row, self.col_date_naissance)
        if date_naissance_str:
            d = self._parse_date(date_naissance_str)
            if d:
                vals['date_naissance'] = d
            else:
                errors.append(f"Date de naissance invalide : '{date_naissance_str}' (format attendu : DD/MM/YYYY)")

        lieu = col(row, self.col_lieu_naissance)
        if lieu:
            vals['lieu_naissance'] = lieu

        sexe_str = col(row, self.col_sexe)
        if sexe_str:
            sexe = self._parse_sexe(sexe_str)
            if sexe:
                vals['sexe'] = sexe
            else:
                errors.append(f"Sexe non reconnu : '{sexe_str}' (M / F attendu)")

        nationalite = col(row, self.col_nationalite)
        vals['nationalite'] = nationalite if nationalite else 'Marocaine'

        adresse = col(row, self.col_adresse)
        if adresse:
            vals['address'] = adresse

        ville = col(row, self.col_ville)
        if ville:
            vals['city'] = ville

        groupe_name = col(row, self.col_groupe)
        vals['groupe_id'] = self._resolve_groupe(groupe_name, dry_run=dry_run)

        contact_nom = col(row, self.col_contact_urgence_nom)
        if contact_nom:
            vals['contact_urgence_nom'] = contact_nom

        contact_tel = col(row, self.col_contact_urgence_tel)
        if contact_tel:
            vals['contact_urgence_tel'] = contact_tel

        return vals, errors

    # ------------------------------------------------------------------
    # Actions wizard
    # ------------------------------------------------------------------

    def action_preview(self):
        """Étape 1 : lire le fichier et afficher un aperçu HTML des 10 premières lignes."""
        self.ensure_one()

        rows = self._parse_file()

        if not rows:
            raise UserError("Le fichier est vide ou ne contient que l'en-tête.")

        # Construire un tableau HTML d'aperçu (10 premières lignes)
        MAX_PREVIEW = 10
        headers = [
            'Ligne', 'Nom complet', 'CNE', 'CIN', 'Email',
            'Téléphone', 'Date naiss.', 'Sexe', 'Groupe', 'Statut'
        ]

        html = ['<table class="table table-sm table-bordered table-striped">']
        html.append('<thead class="thead-dark"><tr>')
        for h in headers:
            html.append(f'<th>{h}</th>')
        html.append('</tr></thead><tbody>')

        for i, row in enumerate(rows[:MAX_PREVIEW], start=1):
            vals, errs = self._build_student_vals(row, dry_run=True)
            status = (
                '<span class="badge badge-danger">❌ Erreur</span>'
                if errs else
                '<span class="badge badge-success">✅ OK</span>'
            )
            cells = [
                str(i),
                self._col(row, self.col_nom),
                self._col(row, self.col_cne),
                self._col(row, self.col_cin),
                self._col(row, self.col_email),
                self._col(row, self.col_telephone),
                self._col(row, self.col_date_naissance),
                self._col(row, self.col_sexe),
                self._col(row, self.col_groupe),
                status + (f'<br/><small class="text-danger">{"; ".join(errs)}</small>' if errs else ''),
            ]
            html.append('<tr>')
            for c in cells:
                html.append(f'<td>{c}</td>')
            html.append('</tr>')

        html.append('</tbody></table>')

        total = len(rows)
        if total > MAX_PREVIEW:
            html.append(
                f'<p class="text-muted">Aperçu limité aux {MAX_PREVIEW} premières lignes '
                f'sur {total} au total.</p>'
            )

        self.write({
            'preview_html': ''.join(html),
            'state': 'preview',
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_import(self):
        """Étape 2 : importer réellement les étudiants."""
        self.ensure_one()

        rows = self._parse_file()

        if not rows:
            raise UserError("Le fichier est vide ou ne contient que l'en-tête.")

        imported = 0
        updated = 0
        skipped = 0
        error_lines = []

        for i, row in enumerate(rows, start=1 + (1 if self.skip_header else 0)):
            vals = {}
            cne_hint = self._col(row, self.col_cne)
            try:
                with self.env.cr.savepoint():
                    vals, errs = self._build_student_vals(row)

                    if errs:
                        error_lines.append(f"Ligne {i} : {' | '.join(errs)}")
                        skipped += 1
                        continue

                    # Vérifier si l'étudiant existe déjà (par CNE)
                    existing = self.env['ensiasd.student'].search(
                        [('cne', '=', vals['cne'])], limit=1
                    )

                    if existing:
                        # Mise à jour partielle (on ne touche pas filiere/niveau si déjà inscrit)
                        update_vals = {k: v for k, v in vals.items()
                                       if k not in ('filiere_id', 'niveau', 'annee_inscription')}
                        existing.write(update_vals)
                        updated += 1
                        _logger.info("Étudiant mis à jour : CNE=%s", vals['cne'])

                    else:
                        # Création
                        student = self.env['ensiasd.student'].create(vals)

                        # Passer à l'état souhaité
                        if self.state_after_import == 'inscrit':
                            student.write({'state': 'inscrit'})
                            if self.auto_inscribe_modules and not student.inscription_ids:
                                student._auto_inscribe_modules()

                        imported += 1
                        _logger.info("Étudiant créé : CNE=%s, Matricule=%s", vals['cne'], student.matricule)

            except ValidationError as ve:
                error_lines.append(f"Ligne {i} (CNE={cne_hint}) : {ve.args[0]}")
                skipped += 1
            except Exception as e:
                error_lines.append(f"Ligne {i} (CNE={cne_hint}) : Erreur inattendue — {str(e)}")
                skipped += 1
                _logger.exception("Erreur import étudiant ligne %d", i)

        summary = (
            f"✅ {imported} étudiant(s) créé(s)  |  "
            f"🔄 {updated} mis à jour  |  "
            f"⚠️ {skipped} ignoré(s) / erreur(s)"
        )

        self.write({
            'import_count': imported,
            'update_count': updated,
            'skip_count': skipped,
            'error_count': len(error_lines),
            'error_log': '\n'.join(error_lines) if error_lines else 'Import terminé sans erreur.',
            'result_summary': summary,
            'state': 'done',
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_download_template_csv(self):
        """Télécharge un fichier CSV modèle pré-rempli avec les en-têtes."""
        header = (
            'nom_complet;cne;cin;email;telephone;mobile;'
            'date_naissance;lieu_naissance;sexe;nationalite;'
            'adresse;ville;groupe;contact_urgence_nom;contact_urgence_tel'
        )
        example = (
            'Mohammed Alami;K123456;AB789012;m.alami@ensiasd.ma;'
            '0612345678;0698765432;15/09/2002;Casablanca;M;Marocaine;'
            '12 rue des roses;Casablanca;G1;Ahmed Alami;0611223344'
        )
        content = f"{header}\n{example}\n"
        file_b64 = base64.b64encode(content.encode('utf-8-sig'))

        attachment = self.env['ir.attachment'].create({
            'name': 'modele_import_etudiants.csv',
            'type': 'binary',
            'datas': file_b64,
            'mimetype': 'text/csv',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_download_template_xlsx(self):
        """Télécharge un fichier Excel modèle avec en-têtes et exemple."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise UserError(
                "Le module 'openpyxl' n'est pas installé.\n"
                "Lancez : pip install openpyxl"
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Étudiants'

        headers = [
            'nom_complet', 'cne', 'cin', 'email', 'telephone', 'mobile',
            'date_naissance', 'lieu_naissance', 'sexe', 'nationalite',
            'adresse', 'ville', 'groupe', 'contact_urgence_nom', 'contact_urgence_tel'
        ]

        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[cell.column_letter].width = max(15, len(h) + 4)

        # Ligne exemple
        example = [
            'Mohammed Alami', 'K123456', 'AB789012', 'm.alami@ensiasd.ma',
            '0612345678', '0698765432', '15/09/2002', 'Casablanca',
            'M', 'Marocaine', '12 rue des roses', 'Casablanca',
            'G1', 'Ahmed Alami', '0611223344'
        ]
        for col_idx, val in enumerate(example, start=1):
            ws.cell(row=2, column=col_idx, value=val)

        # Note d'aide dans un second onglet
        ws2 = wb.create_sheet('Instructions')
        ws2['A1'] = 'Colonne'
        ws2['B1'] = 'Obligatoire'
        ws2['C1'] = 'Format / Valeurs acceptées'
        instructions = [
            ('nom_complet', 'OUI', 'Nom et prénom complets'),
            ('cne', 'OUI', 'Code National Étudiant (unique)'),
            ('cin', 'non', 'Carte Identité Nationale (unique)'),
            ('email', 'non', 'Adresse email valide'),
            ('telephone', 'non', 'Numéro de téléphone fixe'),
            ('mobile', 'non', 'Numéro de téléphone mobile'),
            ('date_naissance', 'non', 'DD/MM/YYYY ou YYYY-MM-DD'),
            ('lieu_naissance', 'non', 'Ville / commune'),
            ('sexe', 'non', 'M ou F  (ou Masculin / Féminin)'),
            ('nationalite', 'non', 'Défaut : Marocaine'),
            ('adresse', 'non', 'Adresse complète'),
            ('ville', 'non', 'Ville de résidence'),
            ('groupe', 'non', 'Nom du groupe (créé automatiquement si absent)'),
            ('contact_urgence_nom', 'non', 'Nom du contact d\'urgence'),
            ('contact_urgence_tel', 'non', 'Téléphone du contact d\'urgence'),
        ]
        for row_idx, (col, req, fmt) in enumerate(instructions, start=2):
            ws2.cell(row=row_idx, column=1, value=col)
            ws2.cell(row=row_idx, column=2, value=req)
            ws2.cell(row=row_idx, column=3, value=fmt)

        buf = io.BytesIO()
        wb.save(buf)
        file_b64 = base64.b64encode(buf.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': 'modele_import_etudiants.xlsx',
            'type': 'binary',
            'datas': file_b64,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_reset(self):
        """Recommencer depuis le début."""
        self.write({
            'state': 'draft',
            'preview_html': False,
            'import_count': 0,
            'update_count': 0,
            'skip_count': 0,
            'error_count': 0,
            'error_log': False,
            'result_summary': False,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }